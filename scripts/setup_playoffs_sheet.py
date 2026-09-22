#!/usr/bin/env python3
"""Actualiza la hoja Playoffs al cuadro de 16 y conserva resultados existentes."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scripts.gsheets import SheetsClient, load_service_account, quote_sheet

ROOT = Path(__file__).resolve().parents[1]
SHEET = "Playoffs"
HEADERS = ["Partida", "Resultado (Jugador,Puntaje,... de 1° a 4°)", "Fecha", "Hora", "Reglas"]
MATCHES = [
    *[(f"QF M{table} G{game}", "A" if table <= 2 else "B") for table in range(1, 5) for game in range(1, 3)],
    *[(f"SF M{table} G{game}", "A" if table == 1 else "B") for table in range(1, 3) for game in range(1, 3)],
    *[(f"FINAL G{game}", "A") for game in range(1, 4)],
]


def cells(previous: dict[str, str] | None = None) -> dict[str, object]:
    previous = previous or {}
    values: dict[str, object] = {
        "A1": "Playoffs — 16 clasificados · Cuartos · Semifinales · Final",
        "A2": "Top 8 de cada división. Avanzan los 2 mejores de cada mesa; cuartos y semifinales tienen 2 hanchan, la final 3.",
        "A4": "Seed A", "B4": "Clasificado División A", "C4": "Seed B", "D4": "Clasificado División B",
        "A15": HEADERS[0], "B15": HEADERS[1], "C15": HEADERS[2], "D15": HEADERS[3], "E15": HEADERS[4],
        "G4": "Mesa", "H4": "Participantes de cuartos", "I4": "Reglas",
    }
    for seed in range(1, 9):
        row = 4 + seed
        values[f"A{row}"] = seed
        values[f"B{row}"] = f"='Clasificación'!B{row}"
        values[f"C{row}"] = seed
        values[f"D{row}"] = f"='Clasificación'!T{row}"
    pairings = [(1, 2, 7, 8), (3, 4, 5, 6), (5, 6, 3, 4), (7, 8, 1, 2)]
    for table, (a1, a2, b1, b2) in enumerate(pairings, start=1):
        row = 4 + table
        values[f"G{row}"] = f"QF M{table}"
        values[f"H{row}"] = f"=B{4+a1}&\", \"&B{4+a2}&\", \"&D{4+b1}&\", \"&D{4+b2}"
        values[f"I{row}"] = f"División {'A' if table <= 2 else 'B'}"
    for offset, (label, division) in enumerate(MATCHES, start=16):
        values[f"A{offset}"] = label
        values[f"B{offset}"] = previous.get(label, "")
        values[f"E{offset}"] = f"División {division}"
    values["A33"] = "Carga oficial"
    values["A34"] = "Pega el resultado ordenado del 1° al 4° en la columna B. Fecha y hora alimentan el calendario del sitio."
    values["A35"] = "El sincronizador calcula el uma según la regla de cada mesa y publica los resultados en Eliminatorias e Historial."
    return values


def update_xlsx(source: Path, destination: Path) -> None:
    workbook = load_workbook(source)
    ws = workbook[SHEET]
    previous = {str(ws.cell(row, 1).value or "").strip().upper(): str(ws.cell(row, 2).value or "").strip() for row in range(1, ws.max_row + 1)}
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
    for address, value in cells(previous).items():
        ws[address] = value
    for row in (1, 4, 15, 33):
        for cell in ws[row][:9]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
    for column, width in {"A": 16, "B": 54, "C": 14, "D": 12, "E": 16, "G": 12, "H": 62, "I": 16}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A16"
    workbook.save(destination)


def update_google(config: dict, credentials: Path | None) -> None:
    client = SheetsClient(config["spreadsheetId"], load_service_account(credentials))
    quoted = quote_sheet(SHEET)
    labels = [f"{quoted}!A{row}" for row in range(1, 101)]
    results = [f"{quoted}!B{row}" for row in range(1, 101)]
    current_labels = client.read_cells(labels)
    current_results = client.read_cells(results)
    previous = {}
    for label_cell, result_cell in zip(labels, results):
        label = current_labels.get(label_cell, "").strip().upper()
        if label:
            previous[label] = current_results.get(result_cell, "")
    updates = {f"{quoted}!{column}{row}": "" for row in range(1, 101) for column in "ABCDEFGHI"}
    updates.update({f"{quoted}!{address}": value for address, value in cells(previous).items()})
    client.write_cells(updates, value_input_option="USER_ENTERED")
    print(f"Hoja {SHEET} actualizada con {client.account_email}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, help="Planilla local que se actualizará")
    parser.add_argument("--output", type=Path, help="Salida; por defecto reemplaza --xlsx")
    parser.add_argument("--write", action="store_true", help="Escribe directamente en Google Sheets")
    parser.add_argument("--credentials", type=Path)
    parser.add_argument("--config", type=Path, default=ROOT / "sync-config.json")
    args = parser.parse_args()
    if args.write:
        update_google(json.loads(args.config.read_text(encoding="utf-8")), args.credentials)
    elif args.xlsx:
        output = args.output or args.xlsx
        update_xlsx(args.xlsx, output)
        print(f"Planilla actualizada: {output}")
    else:
        parser.error("usa --xlsx ARCHIVO o --write")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
