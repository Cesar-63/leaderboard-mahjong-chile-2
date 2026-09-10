import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from scripts.gsheets import quote_sheet
from scripts.majsoul import ParsedPaipu
from scripts.sync import (
    align_history_with_fixtures, format_history_line, history_rows, parse_history,
    parse_history_line, SyncError,
)
from scripts.fill_game_history import (
    apply_writes, build_report, pick_slot, planned_writes, rank_seats, resolve_seat_names,
    score_problems, slots_by_table, slots_in_use, usable_slots, writable_rows,
)
from scripts.fill_calendar_paipus import build_tables


RULE_A = {"uma": [15, 5, -5, -15], "initialPoints": 30000,
          "rosterSheet": "Jugadores Liga A", "historySheet": "Game History A"}
CACHE = Path("data/raw-paipu")
# La celda B69 de Game History A, tal cual la escribe una persona.
EJEMPLO = "MasterFofo,35900,Uznaiker,31100,Mon_96,30900,Tobippi,22100"


def player(pid, name, account_id):
    return {"id": pid, "div": pid[0], "num": pid[1:], "name": name, "shortName": name,
            "handle": name, "accountId": account_id, "discord": "", "nat": "CL"}


ROSTER_A = [
    player("A01", "MasterFofo", 101), player("A02", "Uznaiker", 102),
    player("A03", "Mon_96", 103), player("A04", "Tobippi", 104),
    player("A05", "Kaiser", 105), player("A06", "Nozomi", 106),
    player("A07", "Tanuki", 107), player("A08", "Sora", 108),
]
MESA_1 = ["MasterFofo", "Uznaiker", "Mon_96", "Tobippi"]
MESA_2 = ["Kaiser", "Nozomi", "Tanuki", "Sora"]


def parsed_paipu(seats, scores, uuid="260710-aaaa"):
    """Un `ParsedPaipu` con lo único que mira el script: identidad y puntajes.

    `seats` son pares (account_id, nickname) por asiento; ambos en `None`/"" es
    un registro que no declara quién se sentó dónde.
    """
    return ParsedPaipu(
        uuid=uuid, final_scores=list(scores), hands=8, seat_stats=[],
        players=[{"seat": seat, "account_id": account, "nickname": nickname, "point": None}
                 for seat, (account, nickname) in enumerate(seats)],
        rounds=[], record_game_seen=True, sha256="",
    )


def fixture(session, table, players, division="A"):
    return {"division": division, "session": session, "table": table, "players": list(players),
            "date": "10 jul", "weekday": "vie", "dateISO": "2026-07-10", "time": "20:00"}


def submissions_for(session, table, players, uuids, division="A"):
    """Las dos celdas del Calendario de una mesa, como las emite `read_calendar`."""
    entries = []
    for game in (1, 2):
        uuid = uuids.get(game, "")
        entry = {
            "key": f"{division}-S{session}-M{table}-G{game}", "division": division,
            "session": session, "table": table, "players": list(players), "date": "10 jul",
            "weekday": "vie", "dateISO": "2026-07-10", "time": "20:00", "game": game,
            "cell": f"Calendario!C{10 + game}",
            "url": f"https://mahjongsoul.game.yo-star.com/?paipu={uuid}" if uuid else "",
        }
        if uuid:
            entry["uuid"] = uuid
        entries.append(entry)
    return entries


def history_entry(session, table, game, results, division="A", sheet="Game History A", row=2):
    key = f"{division}-S{session}-M{table}-G{game}"
    return key, {"key": key, "session": session, "table": table, "game": game,
                 "results": results, "sourceCell": f"{sheet}!B{row}"}


def results_from(line, rule=RULE_A):
    return parse_history_line(line, rule, "prueba")


def sheet_rows(sessions=(1,), mesas=(1, 2, 3, 4, 5, 6)):
    """El rótulo de cada fila del Game History, en el orden de la planilla."""
    rows = {}
    row = 2
    for session in sessions:
        for mesa in mesas:
            for game in (1, 2):
                rows[(session, mesa, game)] = row
                row += 1
    return rows


class FormatoDeLaCeldaTests(unittest.TestCase):
    """`Nombre,score,…`: leerlo y escribirlo tienen que ser la misma regla."""

    def test_el_ejemplo_de_la_planilla_va_y_vuelve_igual(self):
        results = results_from(EJEMPLO)
        self.assertEqual([item["name"] for item in results],
                         ["MasterFofo", "Uznaiker", "Mon_96", "Tobippi"])
        self.assertEqual([item["place"] for item in results], [1, 2, 3, 4])
        self.assertEqual(format_history_line(results), EJEMPLO)

    def test_los_puntos_de_liga_salen_del_uma_de_la_division(self):
        results = results_from(EJEMPLO)
        # (35900 - 30000) / 1000 + 15 = 20.9 con el uma de A.
        self.assertEqual(results[0]["delta"], 20.9)
        uma_b = {"uma": [35, 5, -10, -30], "initialPoints": 30000}
        self.assertEqual(results_from(EJEMPLO, uma_b)[0]["delta"], 40.9)

    def test_una_celda_que_no_suma_120000_no_se_acepta(self):
        with self.assertRaises(SyncError) as caught:
            results_from("MasterFofo,35900,Uznaiker,31100,Mon_96,30900,Tobippi,22000")
        self.assertIn("120.000", str(caught.exception))

    def test_espacios_alrededor_de_los_nombres_se_normalizan(self):
        results = results_from(" MasterFofo , 35900 , Uznaiker ,31100, Mon_96,30900,Tobippi,22100")
        self.assertEqual(format_history_line(results), EJEMPLO)

    def test_la_fila_sale_de_la_columna_a_y_no_de_un_calculo(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Game History A"
        ws.cell(1, 1, "Partida")
        ws.cell(1, 2, "Resultado")
        # A propósito con un hueco: el rótulo manda, no la posición.
        ws.cell(2, 1, "S1 M1 G1")
        ws.cell(4, 1, "S1 M1 G2")
        ws.cell(4, 2, EJEMPLO)
        self.assertEqual(history_rows(ws), {(1, 1, 1): 2, (1, 1, 2): 4})
        parsed = parse_history(workbook, "A", "Game History A", RULE_A)
        self.assertEqual(list(parsed), ["A-S1-M1-G2"])
        self.assertEqual(parsed["A-S1-M1-G2"]["sourceCell"], "Game History A!B4")
        self.assertEqual(format_history_line(parsed["A-S1-M1-G2"]["results"]), EJEMPLO)

    def test_el_rango_a1_lleva_la_hoja_entre_comillas(self):
        self.assertEqual(quote_sheet("Game History A") + "!B69", "'Game History A'!B69")


class AsientosYPuestosTests(unittest.TestCase):
    """De los cuatro asientos del paipu a los cuatro puestos de la celda."""

    def test_los_asientos_se_ordenan_de_primero_a_cuarto(self):
        ranked = rank_seats(["Tobippi", "Mon_96", "MasterFofo", "Uznaiker"],
                            [22100, 30900, 35900, 31100])
        self.assertEqual(format_history_line(ranked), EJEMPLO)

    def test_un_empate_lo_desempata_el_asiento_mas_cercano_al_este(self):
        ranked = rank_seats(["MasterFofo", "Uznaiker", "Mon_96", "Tobippi"],
                            [30000, 30000, 30000, 30000])
        self.assertEqual([item["name"] for item in ranked], MESA_1)

    def test_el_paipu_nombra_los_asientos_y_gana_el_nombre_del_roster(self):
        parsed = parsed_paipu([(101, "masterfofo"), (102, "Uznaiker"),
                               (103, "Mon_96"), (104, "Tobippi")],
                              [35900, 31100, 30900, 22100])
        names, identity, _reason = resolve_seat_names(parsed, MESA_1, ROSTER_A)
        # El apodo del juego venía en minúsculas; en la celda va el de la liga.
        self.assertEqual(names, MESA_1)
        self.assertEqual(identity, "paipu")

    def test_un_suplente_entra_con_su_apodo_del_juego(self):
        parsed = parsed_paipu([(101, "MasterFofo"), (102, "Uznaiker"),
                               (103, "Mon_96"), (999, "InvitadoX")],
                              [35900, 31100, 30900, 22100])
        names, identity, _reason = resolve_seat_names(parsed, MESA_1, ROSTER_A)
        self.assertEqual(names, ["MasterFofo", "Uznaiker", "Mon_96", "InvitadoX"])
        self.assertEqual(identity, "paipu")

    def test_sin_identidad_en_el_paipu_el_orden_queda_para_revisar(self):
        parsed = parsed_paipu([(None, ""), (None, ""), (None, ""), (None, "")],
                              [35900, 31100, 30900, 22100])
        names, identity, reason = resolve_seat_names(parsed, MESA_1, ROSTER_A)
        self.assertEqual(names, MESA_1)
        self.assertEqual(identity, "revisar")
        self.assertIn("Calendario", reason)

    def test_sin_paipu_ni_calendario_no_hay_nombres(self):
        parsed = parsed_paipu([(None, ""), (None, ""), (None, ""), (None, "")],
                              [35900, 31100, 30900, 22100])
        names, identity, _reason = resolve_seat_names(parsed, ["", "", "", ""], ROSTER_A)
        self.assertIsNone(names)
        self.assertEqual(identity, "revisar")

    def test_un_paipu_que_no_suma_120000_no_sirve_de_fuente(self):
        self.assertEqual(score_problems([35900, 31100, 30900, 22100]), [])
        problemas = score_problems([35900, 31100, 30900, 22000])
        self.assertEqual(len(problemas), 1)
        self.assertIn("120.000", problemas[0])

    def test_un_puntaje_que_no_es_multiplo_de_100_se_denuncia(self):
        problemas = score_problems([35950, 31050, 30900, 22100])
        self.assertTrue(any("múltiplo de 100" in item for item in problemas))


class MesaDelGameHistoryTests(unittest.TestCase):
    """El Game History numera las mesas distinto: hay que elegir dónde escribir."""

    def test_un_grupo_ya_registrado_conserva_su_mesa_del_game_history(self):
        # El grupo de la mesa 1 del Calendario está en la mesa 4 del historial.
        histories = dict([history_entry(1, 4, 1, results_from(EJEMPLO))])
        fixtures = [fixture(1, 1, MESA_1), fixture(1, 2, MESA_2)]
        aligned, _avisos = align_history_with_fixtures(histories, fixtures)
        self.assertEqual(slots_by_table(aligned, fixtures), {("A", 1, 1): 4})

    def test_un_grupo_que_no_coincide_con_ninguna_mesa_no_reclama_ninguna(self):
        ajenos = results_from("Otro1,35900,Otro2,31100,Otro3,30900,Otro4,22100")
        histories = dict([history_entry(1, 3, 1, ajenos)])
        fixtures = [fixture(1, 1, MESA_1), fixture(1, 2, MESA_2)]
        aligned, _avisos = align_history_with_fixtures(histories, fixtures)
        self.assertEqual(slots_by_table(aligned, fixtures), {})
        # Pero la mesa 3 del historial sigue ocupada y no se puede reutilizar.
        self.assertEqual(slots_in_use(histories), {("A", 1): {3}})

    def test_solo_valen_las_mesas_con_fila_para_los_dos_hanchan(self):
        rows = sheet_rows(mesas=(1, 2))
        del rows[(1, 2, 2)]
        self.assertEqual(usable_slots(rows, 1), {1})

    def test_se_prefiere_el_mismo_numero_que_el_calendario(self):
        self.assertEqual(pick_slot({1, 2, 3}, set(), 2), 2)

    def test_si_el_numero_preferido_esta_ocupado_se_toma_la_primera_libre(self):
        self.assertEqual(pick_slot({1, 2, 3}, {1, 2}, 2), 3)

    def test_sin_mesas_libres_no_se_inventa_ninguna(self):
        self.assertIsNone(pick_slot({1, 2}, {1, 2}, 1))


class ReporteTests(unittest.TestCase):
    """Qué propone el script para cada hanchan del Calendario."""

    def setUp(self):
        self.fixtures = [fixture(1, 1, MESA_1), fixture(1, 2, MESA_2)]
        self.sheets = {"A": {"name": "Game History A", "rows": sheet_rows()}}
        self.rosters = {"A": ROSTER_A}
        self.paipus = {}

    def report(self, submissions, histories=None, slots=None):
        histories = histories or {}
        if slots is None:
            aligned, _avisos = align_history_with_fixtures(histories, self.fixtures)
            slots = slots_by_table(aligned, self.fixtures)
        with patch("scripts.fill_game_history.load_parsed", side_effect=self.fake_load):
            return build_report(build_tables(submissions), self.sheets, self.rosters,
                                histories, slots, CACHE, [])

    def fake_load(self, uuid, _cache_dir):
        entry = self.paipus.get(uuid)
        return entry if entry else (None, "El paipu no está en la caché local; corre con --fetch-logs")

    def mesa_1(self, uuids=("260710-aaaa", "260710-bbbb")):
        for uuid in uuids:
            self.paipus[uuid] = (parsed_paipu(
                [(101, "MasterFofo"), (102, "Uznaiker"), (103, "Mon_96"), (104, "Tobippi")],
                [35900, 31100, 30900, 22100], uuid), None)
        return submissions_for(1, 1, MESA_1, {1: uuids[0], 2: uuids[1]})

    def rows_by_key(self, report):
        return {row["key"]: row for row in report["proposals"]}

    def test_celda_vacia_con_paipu_se_propone_con_su_fila(self):
        report = self.report(self.mesa_1())
        row = self.rows_by_key(report)["A-S1-M1-G1"]
        self.assertEqual(row["status"], "PROPUESTO")
        self.assertEqual(row["value"], EJEMPLO)
        self.assertEqual(row["cell"], "'Game History A'!B2")
        self.assertEqual(row["label"], "S1 M1 G1")
        self.assertEqual(self.rows_by_key(report)["A-S1-M1-G2"]["cell"], "'Game History A'!B3")

    def test_sin_paipu_en_el_calendario_la_celda_queda_pendiente(self):
        report = self.report(submissions_for(1, 2, MESA_2, {}))
        row = self.rows_by_key(report)["A-S1-M2-G1"]
        self.assertEqual(row["status"], "PENDIENTE")
        self.assertIn("paipu", row["detail"])
        self.assertIsNone(row["value"])

    def test_paipu_sin_descargar_no_propone_nada(self):
        submissions = submissions_for(1, 1, MESA_1, {1: "260710-aaaa", 2: "260710-bbbb"})
        report = self.report(submissions)
        row = self.rows_by_key(report)["A-S1-M1-G1"]
        self.assertEqual(row["status"], "PENDIENTE")
        self.assertIn("--fetch-logs", row["detail"])

    def test_una_celda_que_ya_dice_lo_mismo_queda_en_ok(self):
        histories = dict([history_entry(1, 1, 1, results_from(EJEMPLO))])
        report = self.report(self.mesa_1(), histories)
        row = self.rows_by_key(report)["A-S1-M1-G1"]
        self.assertEqual(row["status"], "OK")
        self.assertIsNone(row["value"])

    def test_una_celda_que_dice_otra_cosa_es_conflicto_y_no_se_toca(self):
        otra = "Tobippi,35900,Mon_96,31100,Uznaiker,30900,MasterFofo,22100"
        histories = dict([history_entry(1, 1, 1, results_from(otra))])
        report = self.report(self.mesa_1(), histories)
        rows = self.rows_by_key(report)
        self.assertEqual(rows["A-S1-M1-G1"]["status"], "CONFLICTO")
        self.assertIsNone(rows["A-S1-M1-G1"]["value"])
        # El otro hanchan de la mesa baja a REVISAR: si los dos paipus están
        # cruzados en el Calendario, pegarlo escribiría el resultado del otro.
        self.assertEqual(rows["A-S1-M1-G2"]["status"], "REVISAR")
        self.assertIn("conflicto", rows["A-S1-M1-G2"]["detail"])
        self.assertEqual(writable_rows(report, include_review=False), [])

    def test_el_segundo_hanchan_va_a_la_mesa_que_ya_ocupa_el_grupo(self):
        # G1 del grupo está en la mesa 4 del historial; G2 tiene que ir ahí.
        histories = dict([history_entry(1, 4, 1, results_from(EJEMPLO))])
        report = self.report(self.mesa_1(), histories)
        rows = self.rows_by_key(report)
        self.assertEqual(rows["A-S1-M1-G1"]["status"], "OK")
        self.assertEqual(rows["A-S1-M1-G2"]["status"], "PROPUESTO")
        self.assertEqual(rows["A-S1-M1-G2"]["historyTable"], 4)
        self.assertEqual(rows["A-S1-M1-G2"]["cell"], "'Game History A'!B9")

    def test_una_mesa_nueva_no_pisa_la_mesa_de_otro_grupo(self):
        # La mesa 1 del historial ya es de otro grupo: la del Calendario se corre.
        ajenos = results_from("Kaiser,35900,Nozomi,31100,Tanuki,30900,Sora,22100")
        histories = dict([history_entry(1, 1, 1, ajenos)])
        report = self.report(self.mesa_1(), histories)
        row = self.rows_by_key(report)["A-S1-M1-G1"]
        self.assertEqual(row["status"], "PROPUESTO")
        self.assertEqual(row["historyTable"], 2)

    def test_sin_identidad_en_el_paipu_la_fila_queda_para_revisar(self):
        submissions = submissions_for(1, 1, MESA_1, {1: "260710-aaaa"})
        self.paipus["260710-aaaa"] = (parsed_paipu(
            [(None, ""), (None, ""), (None, ""), (None, "")],
            [35900, 31100, 30900, 22100]), None)
        report = self.report(submissions)
        row = self.rows_by_key(report)["A-S1-M1-G1"]
        self.assertEqual(row["status"], "REVISAR")
        self.assertEqual(row["value"], EJEMPLO)
        self.assertEqual(writable_rows(report, include_review=False), [])
        self.assertEqual(len(writable_rows(report, include_review=True)), 1)

    def test_un_paipu_con_puntajes_imposibles_se_reporta_y_no_se_propone(self):
        submissions = submissions_for(1, 1, MESA_1, {1: "260710-aaaa"})
        self.paipus["260710-aaaa"] = (parsed_paipu(
            [(101, "MasterFofo"), (102, "Uznaiker"), (103, "Mon_96"), (104, "Tobippi")],
            [35900, 31100, 30900, 22000]), None)
        report = self.report(submissions)
        self.assertEqual(self.rows_by_key(report)["A-S1-M1-G1"]["status"], "PENDIENTE")
        self.assertEqual([issue["type"] for issue in report["issues"]], ["PAIPU_INVALIDO"])

    def test_el_resumen_cuenta_cada_estado(self):
        report = self.report(self.mesa_1())
        self.assertEqual(report["summary"]["PROPUESTO"], 2)
        self.assertEqual(report["summary"]["CONFLICTO"], 0)


class FakeSheets:
    """Planilla de mentira: contesta lo que tiene y anota lo que le escriben."""

    def __init__(self, current=None):
        self.current = dict(current or {})
        self.written = {}
        self.spreadsheet_id = "planilla-de-prueba"
        self.account_email = "bot@liga.iam.gserviceaccount.com"

    def read_cells(self, ranges):
        return {cell: self.current.get(cell, "") for cell in ranges}

    def write_cells(self, updates):
        self.current.update(updates)
        self.written.update(updates)
        return len(updates)


class EscrituraEnLaPlanillaTests(unittest.TestCase):
    """`--write`: qué celdas pega el script y cuáles deja para un humano."""

    def filas(self):
        return [{"key": "A-S1-M1-G1", "cell": "'Game History A'!B2", "value": EJEMPLO}]

    def test_pega_la_celda_vacia(self):
        client = FakeSheets()
        resumen = apply_writes(client, self.filas())
        self.assertEqual(resumen["written"], 1)
        self.assertEqual(client.written, {"'Game History A'!B2": EJEMPLO})

    def test_no_pisa_una_celda_que_se_lleno_mientras_tanto(self):
        otra = "Tobippi,35900,Mon_96,31100,Uznaiker,30900,MasterFofo,22100"
        client = FakeSheets({"'Game History A'!B2": otra})
        filas = self.filas()
        resumen = apply_writes(client, filas)
        self.assertEqual(resumen["written"], 0)
        self.assertEqual(resumen["skipped"], 1)
        self.assertEqual(client.written, {})
        self.assertEqual(client.current["'Game History A'!B2"], otra)
        self.assertEqual(filas[0]["writeOutcome"], "OMITIDO")

    def test_una_celda_que_ya_dice_lo_mismo_no_se_reescribe(self):
        client = FakeSheets({"'Game History A'!B2": EJEMPLO.replace(",", ", ")})
        resumen = apply_writes(client, self.filas())
        self.assertEqual((resumen["written"], resumen["already"]), (0, 1))
        self.assertEqual(client.written, {})

    def test_el_simulacro_no_toca_la_planilla(self):
        filas = self.filas()
        resumen = planned_writes(filas)
        self.assertEqual(resumen["mode"], "simulacro")
        self.assertEqual(resumen["attempted"], 1)
        self.assertEqual(filas[0]["writeOutcome"], "SIMULADO")


if __name__ == "__main__":
    unittest.main()
